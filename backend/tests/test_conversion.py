import hashlib
import io
import json
from datetime import timedelta
from urllib.parse import parse_qs, urlparse
from uuid import UUID, uuid4

import httpx
import pytest
from botocore.exceptions import BotoCoreError
from docx import Document
from pypdf import PdfReader
from sqlalchemy import select

from app.api import conversion as routes
from app.conversion import engine
from app.conversion.renderers import ConversionError, render
from app.models.conversion import DocumentConversion
from app.models.documents import Document as StoredDocument
from app.models.documents import DocumentVersion
from app.models.domain import Role, WorkspaceMember
from app.processing import engine as processing
from app.processing.parsers import normalize
from tests.test_documents import complete, put, s3_server, storage, upload  # noqa: F401
from tests.test_foundation import headers
from tests.test_processing import (  # noqa: F401
    archive_bytes,
    docx_bytes,
    pdf_bytes,
    processing_upload,
)

HWPX = archive_bytes(
    {
        "Contents/section0.xml": """<sec xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph"><hp:p><hp:run><hp:t>한글 본문</hp:t></hp:run></hp:p><hp:p><hp:run><hp:tbl><hp:tr><hp:tc><hp:subList><hp:p><hp:run><hp:t>표 내용</hp:t></hp:run></hp:p></hp:subList></hp:tc></hp:tr></hp:tbl></hp:run></hp:p></sec>"""  # noqa: E501
    }
)


@pytest.mark.parametrize("format", ["pdf", "docx", "txt", "md"])
@pytest.mark.parametrize(
    "name,body,expected",
    [
        ("한글.hwpx", HWPX, "한글 본문"),
        ("문서.docx", docx_bytes(), "첫 번째 문단"),
        (
            "note.md",
            (
                "# 한글 제목\n\n한글 본문\n\n- 목록\n\n| 항목 | 값 |\n"
                "| --- | --- |\n| 가 | 나 |\n\n```\nprint('한글')\n```"
            ).encode(),
            "한글 본문",
        ),
        ("note.txt", "한글 본문\n\n둘째 문단".encode(), "한글 본문"),
        ("note.pdf", pdf_bytes(), "Hello PDF"),
    ],
)
def test_five_inputs_to_four_readable_outputs(name, body, expected, format):
    normalized = normalize(body, name, hashlib.sha256(body).hexdigest())
    result, warnings = render(normalized, format)
    assert warnings and result
    if format == "pdf":
        reader = PdfReader(io.BytesIO(result))
        content = "\n".join(page.extract_text() for page in reader.pages)
        assert reader.pages[0]["/Resources"]["/Font"]
    elif format == "docx":
        doc = Document(io.BytesIO(result))
        content = "\n".join(p.text for p in doc.paragraphs)
        if name.endswith((".hwpx", ".docx", ".md")):
            assert doc.tables
    else:
        content = result.decode("utf-8")
    assert expected in content


def test_pdf_long_table_and_korean_text_flow_across_pages():
    rows = [["항목", "설명"]] + [[str(i), "긴 한글 표 내용 " * 30] for i in range(80)]
    doc = {
        "schema_version": 1,
        "blocks": [
            {"type": "heading", "text": "한글 변환 검증", "level": 1},
            {"type": "table", "text": "표", "rows": rows},
            {"type": "paragraph", "text": "끝 문단"},
        ],
    }
    data, _ = render(doc, "pdf")
    reader = PdfReader(io.BytesIO(data))
    assert len(reader.pages) > 2
    assert "끝 문단" in reader.pages[-1].extract_text()
    assert "79" in "".join(p.extract_text() for p in reader.pages)


def test_untrusted_markup_stays_literal_and_code_fence_is_safe():
    literal = '<img src="https://example.test/no-fetch"> [link](javascript:alert(1))'
    doc = {
        "schema_version": 1,
        "blocks": [
            {"type": "paragraph", "text": literal},
            {"type": "code", "text": "```\nprint('hi')\n```"},
        ],
    }
    md, _ = render(doc, "md")
    assert b"\\<img" in md and b"\\[link\\]" in md and b"````\n```" in md
    pdf, _ = render(doc, "pdf")
    reader = PdfReader(io.BytesIO(pdf))
    assert "example.test" in reader.pages[0].extract_text()
    assert not reader.pages[0].get("/Annots")


def test_txt_uses_all_table_cells_and_preserves_code_indentation():
    doc = {
        "schema_version": 1,
        "blocks": [
            {"type": "table", "text": "summary", "rows": [["항목", "값"], ["첫째", "둘째"]]},
            {"type": "code", "text": "def hello():\n    return 1"},
        ],
    }
    body, _ = render(doc, "txt")
    assert body.decode() == "항목\t값\n첫째\t둘째\n\ndef hello():\n    return 1\n"


def test_unsupported_pdf_glyph_and_wide_table_have_explicit_warnings():
    doc = {
        "schema_version": 1,
        "blocks": [
            {"type": "paragraph", "text": "이모지 🦄"},
            {"type": "table", "text": "표", "rows": [[str(i) for i in range(9)]]},
        ],
    }
    body, warnings = render(doc, "pdf")
    assert "?" in PdfReader(io.BytesIO(body)).pages[0].extract_text()
    assert any("글꼴" in w for w in warnings) and any("넓은 표" in w for w in warnings)


@pytest.mark.parametrize(
    "document",
    [
        {},
        {"schema_version": 1, "blocks": []},
        {"schema_version": 1, "blocks": [{"type": "table", "text": "x", "rows": [[3]]}]},
        {"schema_version": 1, "blocks": [{"type": "heading", "text": "x", "level": 99}]},
    ],
)
def test_invalid_normalization_is_rejected(document):
    with pytest.raises(ConversionError):
        render(document, "pdf")


@pytest.fixture
def ready(client, db, processing_upload, storage, monkeypatch):  # noqa: F811
    auth, workspace, ticket = processing_upload
    processing.run(ticket["version_id"])
    monkeypatch.setattr(engine, "get_storage", lambda: storage)
    monkeypatch.setattr(routes, "get_storage", lambda: storage)
    return auth, workspace, ticket


def test_api_converts_all_formats_preserves_original_and_reuses_results(client, db, ready, storage):  # noqa: F811
    auth, _, ticket = ready
    version_id = UUID(ticket["version_id"])
    prefix = f"/api/versions/{version_id}/conversions"
    original = db.get(DocumentVersion, version_id)
    original_hash = original.sha256
    assert client.get(prefix, headers=auth).json() == []
    for format in ("pdf", "docx", "txt", "md"):
        assert client.post(prefix + f"/{format}/download-url", headers=auth).status_code == 409
        first = client.post(prefix + f"/{format}", headers=auth)
        assert first.status_code == 202 and first.json()["state"] == "QUEUED"
        assert first.headers["cache-control"] == "no-store"
        assert client.post(prefix + f"/{format}", headers=auth).json() == first.json()
        engine.run(str(version_id), format)
        engine.run(str(version_id), format)
        db.expire_all()
        job = db.get(DocumentConversion, (version_id, format))
        assert job.state == "COMPLETED" and job.attempts == 1
        assert "/converted/" in job.output_key
        assert client.post(prefix + f"/{format}", headers=auth).json()["state"] == "COMPLETED"
        response = client.post(prefix + f"/{format}/download-url", headers=auth)
        assert response.headers["cache-control"] == "no-store"
        assert response.json()["filename"] == f"노트-v1.{format}"
        url = response.json()["download_url"]
        assert "attachment" in parse_qs(urlparse(url).query)["response-content-disposition"][0]
        body = httpx.get(url)
        assert body.status_code == 200
        assert hashlib.sha256(body.content).hexdigest() == job.sha256
        assert len(body.content) == job.size
        obj = storage.client.head_object(Bucket=original.storage_bucket, Key=job.output_key)
        assert obj["ServerSideEncryption"] == "AES256"
        assert obj["Metadata"]["version-id"] == str(version_id)
    jobs = client.get(prefix, headers=auth).json()
    assert len(jobs) == 4 and all(
        "output_key" not in job and "source_key" not in job for job in jobs
    )
    db.refresh(original)
    assert original.sha256 == original_hash
    assert (
        len(client.get(f"/api/documents/{ticket['document_id']}", headers=auth).json()["versions"])
        == 1
    )


def test_auth_ready_gate_and_viewer_revocation(client, db, processing_upload, monkeypatch, storage):  # noqa: F811
    auth, workspace, ticket = processing_upload
    prefix = f"/api/versions/{ticket['version_id']}/conversions"
    assert client.post(prefix + "/pdf", headers=auth).status_code == 409
    processing.run(ticket["version_id"])
    monkeypatch.setattr(engine, "get_storage", lambda: storage)
    monkeypatch.setattr(routes, "get_storage", lambda: storage)
    viewer, stranger = headers(), headers()
    user = client.get("/api/me", headers=viewer).json()["id"]
    member = WorkspaceMember(workspace_id=UUID(workspace), user_id=UUID(user), role=Role.VIEWER)
    db.add(member)
    db.commit()
    for method, suffix in [("GET", ""), ("POST", "/pdf"), ("POST", "/pdf/download-url")]:
        assert client.request(method, prefix + suffix).status_code == 401
        assert client.request(method, prefix + suffix, headers=stranger).status_code == 404
    assert client.post(prefix + "/exe", headers=auth).status_code == 422
    assert client.post(prefix + "/pdf", headers=viewer).status_code == 202
    engine.run(ticket["version_id"], "pdf")
    assert client.post(prefix + "/pdf/download-url", headers=viewer).status_code == 200
    db.delete(member)
    db.commit()
    assert client.post(prefix + "/pdf/download-url", headers=viewer).status_code == 404


def test_failure_retry_and_source_integrity(client, db, ready, monkeypatch):
    auth, _, ticket = ready
    prefix = f"/api/versions/{ticket['version_id']}/conversions/txt"
    client.post(prefix, headers=auth)
    with monkeypatch.context() as patch:
        patch.setattr(
            engine, "read_object", lambda *args: json.dumps({"source_sha256": "wrong"}).encode()
        )
        engine.run(ticket["version_id"], "txt")
    assert (
        client.get(prefix.rsplit("/", 1)[0], headers=auth).json()[0]["error_code"]
        == "SOURCE_MISMATCH"
    )
    assert client.post(prefix, headers=auth).json()["state"] == "QUEUED"
    engine.run(ticket["version_id"], "txt")
    assert client.get(prefix.rsplit("/", 1)[0], headers=auth).json()[0]["state"] == "COMPLETED"


def test_dispatch_retry_broker_loss_and_expired_lease(client, db, ready, monkeypatch):
    auth, _, ticket = ready
    key = (UUID(ticket["version_id"]), "txt")
    client.post(f"/api/versions/{key[0]}/conversions/txt", headers=auth)
    with pytest.raises(ConnectionError):
        engine.dispatch(lambda *args: (_ for _ in ()).throw(ConnectionError()))
    assert db.get(DocumentConversion, key).dispatched_at is None
    published = []
    engine.dispatch(lambda *args: published.append(args))
    engine.dispatch(lambda *args: published.append(args))
    assert published == [(str(key[0]), "txt")]
    with monkeypatch.context() as patch:
        patch.setattr(engine, "get_storage", lambda: (_ for _ in ()).throw(BotoCoreError()))
        engine.run(str(key[0]), "txt")
    db.expire_all()
    job = db.get(DocumentConversion, key)
    assert job.state == "QUEUED" and job.error_code == "STORAGE_UNAVAILABLE"
    engine.run(str(key[0]), "txt")  # Backoff is enforced even for broker redelivery.
    db.refresh(job)
    assert job.attempts == 1
    job.state, job.started_at, job.attempts = "RUNNING", engine.now() - timedelta(seconds=181), 1
    job.run_token = uuid4()
    db.commit()
    engine.dispatch(lambda *args: published.append(args))
    db.refresh(job)
    assert job.state == "QUEUED" and job.run_token is None
    job.state, job.started_at, job.attempts = "RUNNING", engine.now() - timedelta(seconds=181), 3
    db.commit()
    engine.dispatch(lambda *args: published.append(args))
    db.refresh(job)
    assert job.state == "FAILED" and job.error_code == "WORKER_INTERRUPTED"


def test_stale_worker_cannot_publish_and_jobs_cascade(client, db, ready, monkeypatch):
    auth, _, ticket = ready
    key = (UUID(ticket["version_id"]), "txt")
    client.post(f"/api/versions/{key[0]}/conversions/txt", headers=auth)
    original = engine.render

    def replace_token(*args):
        with engine.SessionLocal() as session:
            session.get(DocumentConversion, key, with_for_update=True).run_token = uuid4()
            session.commit()
        return original(*args)

    monkeypatch.setattr(engine, "render", replace_token)
    engine.run(str(key[0]), "txt")
    job = db.get(DocumentConversion, key)
    assert job.state == "RUNNING" and job.output_key is None
    doc = db.get(StoredDocument, UUID(ticket["document_id"]))
    doc.current_version_id = None
    db.flush()
    db.delete(db.get(DocumentVersion, key[0]))
    db.commit()
    assert db.scalar(select(DocumentConversion)) is None


def test_historical_version_export_is_independent_of_current(client, db, ready):
    auth, _, ticket = ready
    version = db.get(DocumentVersion, UUID(ticket["version_id"]))
    other = DocumentVersion(
        document_id=version.document_id,
        version_number=2,
        storage_bucket=version.storage_bucket,
        storage_key=version.storage_key,
        original_filename="새 이름.md",
        mime_type=version.mime_type,
        size=version.size,
        sha256=version.sha256,
        created_by=version.created_by,
        status=version.status,
        uploaded_at=version.uploaded_at,
    )
    db.add(other)
    db.flush()
    db.get(StoredDocument, version.document_id).current_version_id = other.id
    db.commit()
    path = f"/api/versions/{version.id}/conversions/txt"
    assert client.post(path, headers=auth).status_code == 202
    engine.run(str(version.id), "txt")
    assert client.post(path + "/download-url", headers=auth).json()["filename"] == "노트-v1.txt"
    assert client.get(f"/api/versions/{other.id}/conversions", headers=auth).json() == []


def test_real_celery_conversion_delivery(client, db, ready):
    import time

    from celery.contrib.testing.worker import start_worker

    from app.processing.worker import celery, convert_document

    auth, _, ticket = ready
    key = (UUID(ticket["version_id"]), "pdf")
    client.post(f"/api/versions/{key[0]}/conversions/pdf", headers=auth)
    celery.control.purge()
    with start_worker(celery, perform_ping_check=False, pool="solo", concurrency=1):
        engine.dispatch(lambda version_id, format: convert_document.delay(version_id, format))
        deadline = time.monotonic() + 15
        while time.monotonic() < deadline:
            db.expire_all()
            job = db.get(DocumentConversion, key)
            if job.state == "COMPLETED":
                break
            db.rollback()
            time.sleep(0.1)
        assert job.state == "COMPLETED"
