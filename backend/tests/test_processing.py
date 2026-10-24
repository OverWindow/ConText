import io
import json
import time
import zipfile
from datetime import timedelta
from uuid import UUID, uuid4

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from sqlalchemy import select

from app.api import processing as routes
from app.models.documents import DocumentVersion, VersionStatus
from app.models.domain import Role, WorkspaceMember
from app.models.processing import DocumentProcessing, ProcessingState
from app.processing import engine
from app.processing.parsers import ParseError, normalize
from tests.test_documents import complete, put, s3_server, storage, upload  # noqa: F401
from tests.test_foundation import headers


def pdf_bytes(text=True, encrypted=False):
    writer = PdfWriter()
    page = writer.add_blank_page(612, 792)
    if text:
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
        )
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 72 720 Td (Hello PDF) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
    if encrypted:
        writer.encrypt("test-password")
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def docx_bytes():
    doc = Document()
    doc.add_heading("한글 제목", 1)
    doc.add_paragraph("첫 번째 문단")
    row = doc.add_table(rows=1, cols=2).rows[0]
    row.cells[0].text, row.cells[1].text = "항목", "값"
    doc.add_paragraph("마지막 문단")
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def archive_bytes(files):
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, body in files.items():
            archive.writestr(name, body)
    return output.getvalue()


def test_all_five_parsers_preserve_source_order_and_text():
    docx = normalize(docx_bytes(), "한글.docx", "a" * 64)
    assert [b["type"] for b in docx["blocks"]] == ["heading", "paragraph", "table", "paragraph"]
    assert docx["blocks"][2]["rows"] == [["항목", "값"]]
    assert docx["title"] == "한글 제목"
    pdf = normalize(pdf_bytes(), "paper.pdf", "a" * 64)
    assert pdf["blocks"][0]["text"] == "Hello PDF"
    assert pdf["blocks"][0]["source"] == {"page": 1}
    xml = '<hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph"><hp:p><hp:run><hp:t>한글 본문</hp:t></hp:run></hp:p><hp:p><hp:run><hp:tbl><hp:tr><hp:tc><hp:subList><hp:p><hp:run><hp:t>표 내용</hp:t></hp:run></hp:p></hp:subList></hp:tc></hp:tr></hp:tbl></hp:run></hp:p></hs:sec>'  # noqa: E501
    hwpx = normalize(archive_bytes({"Contents/section0.xml": xml}), "문서.hwpx", "b" * 64)
    assert [b["text"] for b in hwpx["blocks"]] == ["한글 본문", "표 내용"]
    assert hwpx["blocks"][1]["type"] == "table"
    md = normalize(
        "# 제목\n\n문단 **강조**\n\n- 목록\n\n| 키 | 값 |\n| --- | --- |\n| a | b |\n\n```js\nconst a = 1;\n```".encode(),  # noqa: E501
        "note.md",
        "c" * 64,
    )
    assert [b["type"] for b in md["blocks"]] == [
        "heading",
        "paragraph",
        "list_item",
        "table",
        "code",
    ]
    assert md["blocks"][1]["text"] == "문단 강조"
    assert md["source_sha256"] == "c" * 64
    for encoding in ["utf-8-sig", "utf-16", "cp949"]:
        txt = normalize("첫 문단\r\n\r\n둘째 문단".encode(encoding), "메모.txt", "d" * 64)
        assert len(txt["blocks"]) == 2
        assert txt["blocks"][1]["text"] == "둘째 문단"


@pytest.mark.parametrize(
    "body,name,code",
    [
        (b"broken", "bad.pdf", "INVALID_FILE"),
        (pdf_bytes(False), "scan.pdf", "OCR_REQUIRED"),
        (pdf_bytes(encrypted=True), "locked.pdf", "ENCRYPTED_PDF"),
        (b"\x00binary", "bad.txt", "INVALID_TEXT"),
        (b"", "empty.md", "EMPTY_DOCUMENT"),
        (archive_bytes({"../danger.xml": "<x/>"}), "bad.hwpx", "INVALID_ARCHIVE"),
        (
            archive_bytes(
                {
                    "Contents/section0.xml": '<!DOCTYPE x [<!ENTITY x SYSTEM "file:///etc/passwd">]><x>&x;</x>'
                }
            ),
            "bad.hwpx",
            "INVALID_ARCHIVE",
        ),
        (archive_bytes({"Contents/section0.xml": "a" * 1000000}), "bomb.hwpx", "LIMIT_EXCEEDED"),
        (b"a" * 2000001, "huge.txt", "LIMIT_EXCEEDED"),
        (b"hello", "bad.exe", "UNSUPPORTED_FORMAT"),
    ],
)
def test_malformed_encrypted_empty_and_bounded_input(body, name, code):
    with pytest.raises(ParseError) as exc:
        normalize(body, name, "a" * 64)
    assert exc.value.code == code


@pytest.fixture
def processing_upload(client, upload, storage, monkeypatch):  # noqa: F811
    monkeypatch.setattr(engine, "get_storage", lambda: storage)
    monkeypatch.setattr(routes, "get_storage", lambda: storage)
    auth, workspace, ticket = upload
    assert put(ticket).status_code == 200
    assert complete(client, auth, ticket).status_code == 200
    return auth, workspace, ticket


def test_processing_artifacts_idempotence_and_pagination(client, db, storage, processing_upload):  # noqa: F811
    auth, _, ticket = processing_upload
    version_id = UUID(ticket["version_id"])
    assert db.get(DocumentProcessing, version_id).state == ProcessingState.QUEUED
    engine.run(str(version_id))
    engine.run(str(version_id))
    db.expire_all()
    job = db.get(DocumentProcessing, version_id)
    assert job.state == ProcessingState.COMPLETED and job.attempts == 1
    version = db.get(DocumentVersion, version_id)
    assert version.status == VersionStatus.READY
    assert "/processed/" in job.normalized_key
    artifact = storage.client.get_object(Bucket=version.storage_bucket, Key=job.normalized_key)
    assert artifact["ServerSideEncryption"] == "AES256"
    assert json.loads(artifact["Body"].read())["source_sha256"] == version.sha256
    prefix = f"/api/versions/{version_id}"
    result = client.get(prefix + "/processing", headers=auth)
    assert "normalized_key" not in result.json()
    assert result.json()["summary"]["block_count"] == 1
    content = client.get(prefix + "/content?limit=1", headers=auth)
    assert content.status_code == 200
    assert content.headers["cache-control"] == "no-store"
    assert content.json()["blocks"][0]["text"] == "Phase 2 document"
    assert client.get(prefix + "/content?offset=1", headers=auth).json()["blocks"] == []
    assert client.get(prefix + "/content?limit=201", headers=auth).status_code == 422
    assert client.post(prefix + "/process", headers=auth).json()["state"] == "COMPLETED"
    assert complete(client, auth, ticket).json()["versions"][0]["status"] == "READY"


def test_failure_retry_preserves_original_and_authorization(
    client, db, processing_upload, monkeypatch
):
    auth, workspace, ticket = processing_upload
    prefix = f"/api/versions/{ticket['version_id']}"

    def fail(*args):
        raise ParseError("OCR_REQUIRED", "스캔 PDF는 OCR 처리가 필요합니다.")

    with monkeypatch.context() as patch:
        patch.setattr(engine, "normalize", fail)
        engine.run(ticket["version_id"])
    assert client.get(prefix + "/processing", headers=auth).json()["error_code"] == "OCR_REQUIRED"
    assert client.get(prefix + "/content", headers=auth).status_code == 409
    assert client.post(prefix + "/download-url", headers=auth).status_code == 200
    viewer, stranger = headers(), headers()
    user = client.get("/api/me", headers=viewer).json()["id"]
    db.add(WorkspaceMember(workspace_id=UUID(workspace), user_id=UUID(user), role=Role.VIEWER))
    db.commit()
    assert client.get(prefix + "/processing", headers=viewer).status_code == 200
    assert client.post(prefix + "/process", headers=viewer).status_code == 403
    for method, endpoint in [("GET", "processing"), ("GET", "content"), ("POST", "process")]:
        assert client.request(method, prefix + "/" + endpoint, headers=stranger).status_code == 404
        assert client.request(method, prefix + "/" + endpoint).status_code == 401
    assert client.post(prefix + "/process", headers=auth).status_code == 202
    engine.run(ticket["version_id"])
    assert client.get(prefix + "/content", headers=viewer).status_code == 200
    members = client.get(f"/api/workspaces/{workspace}/members", headers=auth)
    assert len(members.json()) == 2
    assert all("email" not in item for item in members.json())
    assert client.get(f"/api/workspaces/{workspace}/members", headers=stranger).status_code == 404


def test_dispatch_recovers_broker_outage_and_stale_lease(db, processing_upload):
    _, _, ticket = processing_upload

    def unavailable(version_id):
        raise ConnectionError("broker unavailable")

    with pytest.raises(ConnectionError):
        engine.dispatch(unavailable)
    job = db.get(DocumentProcessing, UUID(ticket["version_id"]))
    assert job.dispatched_at is None
    published = []
    engine.dispatch(published.append)
    engine.dispatch(published.append)
    assert published == [ticket["version_id"]]
    db.refresh(job)
    job.state = ProcessingState.RUNNING
    job.started_at = engine.now() - timedelta(seconds=181)
    job.attempts = 1
    job.run_token = uuid4()
    db.commit()
    engine.dispatch(published.append)
    db.refresh(job)
    assert job.state == ProcessingState.QUEUED and job.run_token is None
    assert len(published) == 2
    job.state = ProcessingState.RUNNING
    job.started_at = engine.now() - timedelta(seconds=181)
    job.attempts = 3
    db.commit()
    engine.dispatch(published.append)
    db.refresh(job)
    assert job.state == ProcessingState.FAILED


def test_stale_worker_cannot_commit_result(db, processing_upload, monkeypatch):
    _, _, ticket = processing_upload
    original = engine.normalize

    def replace_lease(*args):
        with engine.SessionLocal() as session:
            job = session.get(DocumentProcessing, UUID(ticket["version_id"]))
            job.run_token = uuid4()
            session.commit()
        return original(*args)

    monkeypatch.setattr(engine, "normalize", replace_lease)
    engine.run(ticket["version_id"])
    db.expire_all()
    job = db.get(DocumentProcessing, UUID(ticket["version_id"]))
    assert job.normalized_key is None and job.state == ProcessingState.RUNNING


def test_source_checksum_changed_is_rejected(db, storage, processing_upload):  # noqa: F811
    _, _, ticket = processing_upload
    version = db.get(DocumentVersion, UUID(ticket["version_id"]))
    storage.client.put_object(
        Bucket=version.storage_bucket, Key=version.storage_key, Body=b"changed"
    )
    engine.run(ticket["version_id"])
    db.expire_all()
    assert db.get(DocumentProcessing, version.id).error_code == "SOURCE_MISMATCH"


def test_real_redis_celery_delivery(db, processing_upload):
    from celery.contrib.testing.worker import start_worker

    from app.processing.worker import celery, process_document

    _, _, ticket = processing_upload
    celery.control.purge()
    with start_worker(celery, perform_ping_check=False, pool="solo", concurrency=1):
        engine.dispatch(lambda version_id: process_document.delay(version_id))
        deadline = time.monotonic() + 15
        while time.monotonic() < deadline:
            db.expire_all()
            job = db.scalar(
                select(DocumentProcessing).where(
                    DocumentProcessing.version_id == UUID(ticket["version_id"])
                )
            )
            if job.state == ProcessingState.COMPLETED:
                break
            db.rollback()
            time.sleep(0.1)
        assert job.state == ProcessingState.COMPLETED
