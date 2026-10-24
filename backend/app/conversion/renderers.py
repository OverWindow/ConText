"""Render only normalized text/structure; never load external links, images or HTML."""

import io
import re
import threading
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.platypus.doctemplate import LayoutError

MIME = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "txt": "text/plain; charset=utf-8",
    "md": "text/markdown; charset=utf-8",
}
FIDELITY = (
    "텍스트·제목·목록·표 중심으로 변환했습니다. 이미지·원본 배치·복잡한 서식은 포함되지 않습니다."
)
FONT_PATH = Path("/usr/share/fonts/truetype/nanum/NanumGothic.ttf")
_font_lock = threading.Lock()


class ConversionError(Exception):
    def __init__(self, code, message):
        super().__init__(message)
        self.code, self.message = code, message


def validate(document):
    if not isinstance(document, dict) or document.get("schema_version") != 1:
        raise ConversionError("INVALID_NORMALIZED", "정규화 문서 형식을 읽을 수 없습니다.")
    blocks = document.get("blocks")
    if not isinstance(blocks, list) or not blocks or len(blocks) > 20000:
        raise ConversionError("INVALID_NORMALIZED", "변환할 문단이 없거나 문단 수가 너무 많습니다.")
    total = 0
    cell_count = 0
    for block in blocks:
        if (
            not isinstance(block, dict)
            or block.get("type") not in ("heading", "paragraph", "list_item", "table", "code")
            or not isinstance(block.get("text"), str)
        ):
            raise ConversionError("INVALID_NORMALIZED", "지원하지 않는 문단 구조입니다.")
        total += len(block["text"])
        if block["type"] == "heading" and (
            type(block.get("level", 1)) is not int or not 1 <= block.get("level", 1) <= 6
        ):
            raise ConversionError("INVALID_NORMALIZED", "제목 단계가 올바르지 않습니다.")
        if block["type"] == "table":
            rows = block.get("rows")
            if (
                not isinstance(rows, list)
                or not rows
                or len(rows) > 10000
                or any(
                    not isinstance(row, list)
                    or not row
                    or len(row) > 100
                    or any(not isinstance(cell, str) for cell in row)
                    for row in rows
                )
            ):
                raise ConversionError(
                    "INVALID_NORMALIZED", "표 구조가 올바르지 않거나 너무 큽니다."
                )
            total += sum(len(cell) for row in rows for cell in row)
            cell_count += sum(map(len, rows))
    if total > 4_000_000 or cell_count > 50000:
        raise ConversionError("LIMIT_EXCEEDED", "변환 문서의 텍스트가 처리 한도를 초과했습니다.")
    return blocks


def clean(text):
    # XML 1.0, ReportLab and Word require valid scalar characters.
    return "".join(
        c
        for c in text
        if c in "\n\t"
        or 0x20 <= ord(c) <= 0xD7FF
        or 0xE000 <= ord(c) <= 0xFFFD
        or 0x10000 <= ord(c) <= 0x10FFFF
    )


def markdown_text(text):
    return re.sub(r"([\\`*_{}\[\]()<>#+.!|~\-])", r"\\\1", clean(text))


def render_text(blocks, markdown=False):
    parts = []
    for b in blocks:
        text = clean(b["text"])
        if not markdown:
            if b["type"] == "table":
                text = "\n".join("\t".join(clean(cell) for cell in row) for row in b["rows"])
            parts.append(text)
        elif b["type"] == "heading":
            parts.append("#" * b.get("level", 1) + " " + markdown_text(text).replace("\n", " "))
        elif b["type"] == "list_item":
            parts.append("- " + markdown_text(text).replace("\n", "\n  "))
        elif b["type"] == "code":
            fence = "`" * max(3, max((len(m[0]) + 1 for m in re.finditer(r"`+", text)), default=3))
            parts.append(f"{fence}\n{text}\n{fence}")
        elif b["type"] == "table":
            rows = b["rows"]
            width = max(map(len, rows))
            lines = [
                "| "
                + " | ".join(
                    markdown_text(cell).replace("\n", "<br>")
                    for cell in row + [""] * (width - len(row))
                )
                + " |"
                for row in rows
            ]
            lines.insert(1, "| " + " | ".join(["---"] * width) + " |")
            parts.append("\n".join(lines))
        else:
            parts.append(markdown_text(text).replace("\n", "  \n"))
    return ("\n\n".join(parts) + "\n").encode("utf-8")


def render_docx(blocks):
    doc = Document()
    doc.core_properties.title = "ConText 문서 내보내기"
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.7)
    for name in [
        "Normal",
        "Title",
        "Heading 1",
        "Heading 2",
        "Heading 3",
        "Heading 4",
        "Heading 5",
        "Heading 6",
        "List Bullet",
        "No Spacing",
    ]:
        style = doc.styles[name]
        style.font.name = "NanumGothic"
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "나눔고딕")
    doc.styles["Normal"].font.size = Pt(10.5)
    for b in blocks:
        if b["type"] == "table":
            width = max(map(len, b["rows"]))
            table = doc.add_table(rows=0, cols=width)
            table.style = "Table Grid"
            for cells in b["rows"]:
                row = table.add_row()
                for cell, value in zip(row.cells, cells):
                    cell.text = clean(value)
            # Repeat the first row across page breaks.
            header = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(header)
        elif b["type"] == "heading":
            doc.add_heading(clean(b["text"]), b.get("level", 1))
        else:
            style = (
                "List Bullet"
                if b["type"] == "list_item"
                else "No Spacing"
                if b["type"] == "code"
                else "Normal"
            )
            doc.add_paragraph(clean(b["text"]), style)
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def render_pdf(blocks, warnings):
    with _font_lock:
        if "ConTextKorean" not in pdfmetrics.getRegisteredFontNames():
            if not FONT_PATH.exists():
                raise ConversionError("FONT_UNAVAILABLE", "PDF 한글 글꼴을 사용할 수 없습니다.")
            pdfmetrics.registerFont(TTFont("ConTextKorean", str(FONT_PATH)))
    font = pdfmetrics.getFont("ConTextKorean")
    unsupported = False

    def text(value):
        nonlocal unsupported
        result = []
        for c in clean(value):
            if c not in "\n\t" and ord(c) not in font.face.charToGlyph:
                unsupported = True
                c = "?"
            result.append(c)
        return escape("".join(result)).replace("\t", "    ").replace("\n", "<br/>")

    body = ParagraphStyle(
        "body",
        fontName="ConTextKorean",
        fontSize=10.5,
        leading=16,
        wordWrap="CJK",
        spaceAfter=9,
        alignment=TA_LEFT,
    )
    cell_style = ParagraphStyle("cell", parent=body, fontSize=9, leading=13, spaceAfter=0)
    story = []
    for b in blocks:
        if b["type"] == "table":
            width = max(map(len, b["rows"]))
            if width > 8:
                if "넓은 표는 행별 텍스트로 펼쳐 표시했습니다." not in warnings:
                    warnings.append("넓은 표는 행별 텍스트로 펼쳐 표시했습니다.")
                for cells in b["rows"]:
                    story.append(
                        Paragraph(
                            "<br/>".join(f"{i + 1}. {text(v)}" for i, v in enumerate(cells)), body
                        )
                    )
                continue
            rows = [
                [Paragraph(text(v), cell_style) for v in row + [""] * (width - len(row))]
                for row in b["rows"]
            ]
            table = Table(
                rows,
                colWidths=[(A4[0] - 108) / width] * width,
                repeatRows=1,
                splitByRow=1,
                splitInRow=1,
                hAlign="LEFT",
            )
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EDF0FF")),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9E1F2")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 7),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                        ("TOPPADDING", (0, 0), (-1, -1), 6),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ]
                )
            )
            story.extend([table, Spacer(1, 12)])
        else:
            style = body
            if b["type"] == "heading":
                size = max(12, 23 - b.get("level", 1) * 2)
                style = ParagraphStyle(
                    "heading",
                    parent=body,
                    fontSize=size,
                    leading=size * 1.4,
                    textColor=colors.HexColor("#334155"),
                    spaceBefore=12,
                )
            elif b["type"] == "code":
                style = ParagraphStyle(
                    "code",
                    parent=body,
                    fontSize=9,
                    leading=14,
                    backColor=colors.HexColor("#F7F9FC"),
                    borderPadding=5,
                )
            content = ("• " if b["type"] == "list_item" else "") + b["text"]
            markup = text(content)
            if b["type"] == "code":
                markup = markup.replace(" ", "&nbsp;")
            story.append(Paragraph(markup, style))

    def footer(canvas, doc):
        if doc.page > 500:
            raise ConversionError("LIMIT_EXCEEDED", "변환 PDF는 최대 500페이지까지 지원합니다.")
        canvas.saveState()
        canvas.setFont("ConTextKorean", 8)
        canvas.setFillColor(colors.HexColor("#64748B"))
        canvas.drawRightString(A4[0] - 48, 25, str(doc.page))
        canvas.restoreState()

    output = io.BytesIO()
    try:
        SimpleDocTemplate(
            output,
            pagesize=A4,
            leftMargin=48,
            rightMargin=48,
            topMargin=42,
            bottomMargin=42,
            title="ConText 문서 내보내기",
            author="ConText",
        ).build(story, onFirstPage=footer, onLaterPages=footer)
    except LayoutError as exc:
        raise ConversionError(
            "LAYOUT_LIMIT", "PDF에 배치하기 어려운 표입니다. DOCX 또는 TXT로 내보내 주세요."
        ) from exc
    if unsupported:
        warnings.append(
            "PDF 글꼴에 없는 일부 문자를 ?로 표시했습니다. "
            "원문 보존은 TXT 또는 DOCX를 사용해 주세요."
        )
    return output.getvalue()


def render(document, format):
    if format not in MIME:
        raise ConversionError("UNSUPPORTED_FORMAT", "지원하지 않는 출력 형식입니다.")
    blocks = validate(document)
    warnings = [FIDELITY] + [w[:500] for w in document.get("warnings", []) if isinstance(w, str)][
        :20
    ]
    if format in ("txt", "md"):
        body = render_text(blocks, markdown=format == "md")
    elif format == "docx":
        body = render_docx(blocks)
    else:
        body = render_pdf(blocks, warnings)
    if len(body) > 50 * 1024 * 1024:
        raise ConversionError("LIMIT_EXCEEDED", "변환 결과가 50MB 한도를 초과했습니다.")
    return body, warnings
